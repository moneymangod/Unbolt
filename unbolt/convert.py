"""Format registry: maps a locked file extension to the outputs it can be converted to."""
import io
import os
import re

from pptx import Presentation
from pptx.util import Emu

from unbolt import daw
from unbolt.als import load_als
from unbolt.aup3 import load_aup3
from unbolt.figfile import render_fig
from unbolt.flp import load_flp

IMAGES, SONGS, AUDIO = ("png", "pdf", "pptx"), ("mp3", "mp4", "wav"), ("wav", "mp3")
KIND_LABELS = {"image": "Designs", "song": "Music projects", "audio": "Audio projects"}
# To support another format, add: extension -> (output formats, loader, kind).
# Loaders return: image -> (frames, warnings); song -> Song; audio -> (samples, warnings).
FORMATS = {
    ".fig": (IMAGES, render_fig, "image"),
    ".als": (SONGS, load_als, "song"),
    ".flp": (SONGS, load_flp, "song"),
    ".aup3": (AUDIO, load_aup3, "audio"),
}


def _safe(name):
    return re.sub(r'[<>:"/\\|?*]+', "_", name).strip() or "output"


def _free(path):
    """Return `path`, or 'name (2).ext' etc. if it exists, so a conversion never overwrites a file."""
    base, ext = os.path.splitext(path)
    n, candidate = 1, path
    while os.path.exists(candidate):
        n += 1
        candidate = f"{base} ({n}){ext}"
    return candidate


def _pptx(frames, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for name, img in frames:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        k = min(prs.slide_width / img.width, prs.slide_height / img.height)
        w, h = int(img.width * k), int(img.height * k)
        buf = io.BytesIO()
        img.save(buf, "PNG")
        slide.shapes.add_picture(buf, (prs.slide_width - w) // 2, (prs.slide_height - h) // 2, w, h)
        slide.notes_slide.notes_text_frame.text = name
    prs.save(path)


def convert(path, fmt, outdir):
    """Convert `path` to `fmt` in `outdir`. Returns (output files, warnings)."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in FORMATS:
        raise ValueError(f"Unsupported file type: {ext or '(none)'}")
    outputs, loader, kind = FORMATS[ext]
    if fmt not in outputs:
        raise ValueError(f"{ext} files can be converted to: {', '.join(outputs)}")
    os.makedirs(outdir, exist_ok=True)
    stem = _safe(os.path.splitext(os.path.basename(path))[0])
    out = _free(os.path.join(outdir, f"{stem}.{fmt}"))
    result = loader(path)
    if kind == "song":
        daw.export(result, out, fmt)
        return [out], result.warnings
    if kind == "audio":
        audio, warnings = result
        daw.write_audio(audio, out, fmt)
        return [out], warnings
    frames, warnings = result
    if fmt == "png":
        files = []
        for i, (name, img) in enumerate(frames, 1):
            files.append(_free(os.path.join(outdir, f"{stem} - {i:02d} {_safe(name)}.png")))
            img.save(files[-1])
        return files, warnings
    if fmt == "pdf":
        frames[0][1].save(out, save_all=True, append_images=[i for _, i in frames[1:]])
    else:
        _pptx(frames, out)
    return [out], warnings
